import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPT_PATH = ROOT / "scripts" / "create_erasmus_budget_workbook.py"


class ErasmusBudgetWorkbookTests(unittest.TestCase):
    def load_module(self):
        if not SCRIPT_PATH.exists():
            self.fail(f"Generator script missing: {SCRIPT_PATH}")
        spec = importlib.util.spec_from_file_location("erasmus_budget_workbook", SCRIPT_PATH)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        return module

    def test_generates_expected_workbook_structure_and_totals(self):
        module = self.load_module()

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "budget.xlsx"
            module.build_budget_workbook(output_path)

            wb = load_workbook(output_path, data_only=False)
            self.assertEqual(wb.sheetnames, ["Overview", "Partner_Detail", "Checks"])

            overview = wb["Overview"]
            detail = wb["Partner_Detail"]
            checks = wb["Checks"]

            self.assertEqual(overview["A1"].value, "Opening: The Granular Financial Breakdown")
            self.assertEqual(overview["A5"].value, "Partner")
            self.assertEqual(overview["A13"].value, "Consortium Total")
            self.assertEqual(overview["H13"].value, "=SUM(H6:H12)")

            self.assertEqual(detail["A4"].value, "Partner")
            self.assertEqual(checks["A4"].value, "Partner")

            expected_partners = [
                "UNIVERSITY OF MALTA (UM)",
                "HUMBOLDT-UNIVERSITÄT ZU BERLIN (UBER)",
                "ELLINOGERMANIKI AGOGI (EA)",
                "MALTA FURTHER AND HIGHER EDUCATION AUTHORITY (MFHEA)",
                "CHAMBER OF COMMERCE, INDUSTRY AND SERVICES OF TERRASSA (CCIT)",
                "T2 S.R.L.",
                "LITHUANIAN UNIVERSITY OF HEALTH SCIENCES (LSMU)",
            ]
            for row_index, partner in enumerate(expected_partners, start=6):
                self.assertEqual(overview[f"A{row_index}"].value, partner)

            self.assertEqual(overview["H6"].value, "=SUM(B6:G6)")
            self.assertEqual(overview["H12"].value, "=SUM(B12:G12)")
            self.assertEqual(overview["H13"].value, "=SUM(H6:H12)")

            stated_totals = [235000, 196000, 155000, 136000, 180000, 160000, 188000]
            for row_index, stated_total in enumerate(stated_totals, start=5):
                self.assertEqual(checks[f"B{row_index}"].value, stated_total)
                self.assertTrue(str(checks[f"C{row_index}"].value).startswith("="))
                self.assertEqual(checks[f"D{row_index}"].value, f"=B{row_index}-C{row_index}")
                self.assertEqual(checks[f"G{row_index}"].value, f'=IF(ABS(D{row_index})<0.01,"OK","CHECK")')

            self.assertEqual(checks["B13"].value, 1250000)
            self.assertEqual(checks["D13"].value, "=B13-C13")
            self.assertEqual(checks["G13"].value, '=IF(ABS(D13)<0.01,"OK","CHECK")')


if __name__ == "__main__":
    unittest.main()
